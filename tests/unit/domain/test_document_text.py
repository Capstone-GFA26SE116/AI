"""FR-08b-2: file type detection and text extraction (T2-T6)."""

import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from metanoia_ai.domain.document_text import (
    DocumentError,
    DocumentErrorCode,
    DocumentKind,
    detect_kind,
    extension_of,
    extract_text,
    normalize_text,
)

MB = 1024 * 1024
PNG = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
BRIEF_LINES = [
    "Website ban hang cho tiem banh Mai.",
    "Can trang chu, trang san pham va gio hang.",
    "Han chot cuoi thang 11.",
]


def make_pdf(pages: list[list[str]], password: str | None = None) -> bytes:
    """A real PDF with text on each page (Helvetica, ASCII only)."""
    writer = PdfWriter()
    font = writer._add_object(
        DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
    )
    for lines in pages:
        page = writer.add_blank_page(612, 792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        shown = " T* ".join(f"({line}) Tj" for line in lines)
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 14 TL 72 720 Td {shown} ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    if password:
        writer.encrypt(user_password=password, algorithm="RC4-128")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Yeu cau du an")
    document.add_paragraph("")
    document.add_paragraph("")
    document.add_paragraph("Pham vi cong viec")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Hang muc"
    table.cell(0, 1).text = "Thoi gian"
    table.cell(1, 0).text = "Thiet ke"
    table.cell(1, 1).text = "5 ngay"
    document.add_paragraph("Ket thuc")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_zip() -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("notes.txt", "not a word file")
    return out.getvalue()


def read(data: bytes, path: str = "brief.pdf", **kw) -> str:
    options = {"max_bytes": 20 * MB, "max_pdf_pages": 100} | kw
    return extract_text(data, name="Brief", path=path, **options)


def error_of(data: bytes, path: str, **kw) -> DocumentError:
    with pytest.raises(DocumentError) as caught:
        read(data, path, **kw)
    return caught.value


# --- T2: PDF with text ---


def test_t2_pdf_text_is_read_from_every_page():
    data = make_pdf([BRIEF_LINES[:2], BRIEF_LINES[2:]])
    text = read(data)
    for line in BRIEF_LINES:
        assert line in text
    assert text.index(BRIEF_LINES[0]) < text.index(BRIEF_LINES[2])


def test_t2_pdf_stops_at_the_page_limit():
    data = make_pdf([[f"Page number {n} has enough words to count."] for n in range(5)])
    text = read(data, max_pdf_pages=2)
    assert "Page number 1" in text
    assert "Page number 2" not in text


def test_t2_pdf_stops_once_it_has_enough_text():
    data = make_pdf([[f"Page number {n} has enough words to count."] for n in range(5)])
    text = read(data, enough_chars=60)
    assert "Page number 1" in text
    assert "Page number 2" not in text


# --- T3: PDF without text, locked PDF, broken PDF ---


def test_t3_pdf_without_text():
    assert error_of(make_pdf([[]]), "scan.pdf").code is DocumentErrorCode.NO_TEXT


def test_t3_pdf_with_too_little_text():
    error = error_of(make_pdf([["Short"]]), "scan.pdf")
    assert error.code is DocumentErrorCode.NO_TEXT


def test_t3_locked_pdf():
    data = make_pdf([BRIEF_LINES], password="secret")
    error = error_of(data, "locked.pdf")
    assert error.code is DocumentErrorCode.UNREADABLE
    assert error.name == "Brief"


def test_t3_broken_pdf():
    error = error_of(b"%PDF-1.7 this is not really a pdf", "broken.pdf")
    assert error.code is DocumentErrorCode.UNREADABLE


# --- T4: DOCX with paragraphs and a table ---


def test_t4_docx_reads_paragraphs_and_table_cells():
    text = read(make_docx(), "yeu-cau.docx")
    assert "Yeu cau du an" in text
    assert "Hang muc | Thoi gian" in text
    assert "Thiet ke | 5 ngay" in text
    assert text.index("Pham vi cong viec") < text.index("Hang muc")
    assert text.index("5 ngay") < text.index("Ket thuc")
    assert "\n\n\n" not in text  # empty paragraphs collapsed


def test_t4_broken_docx():
    data = make_docx()
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.namelist():
            body = b"<broken" if item == "word/document.xml" else src.read(item)
            dst.writestr(item, body)
    error = error_of(out.getvalue(), "broken.docx")
    assert error.code is DocumentErrorCode.UNREADABLE


# --- T5: text files and files we do not read ---


@pytest.mark.parametrize("path", ["notes.txt", "README.MD", "uploads/a/b.md"])
def test_t5_utf8_text_files(path):
    data = "\ufeffYêu cầu:\r\n- Trang chủ\r\n\r\n\r\n\r\n- Giỏ hàng  \n".encode()
    assert read(data, path) == "Yêu cầu:\n- Trang chủ\n\n- Giỏ hàng"


@pytest.mark.parametrize(
    ("data", "path"),
    [
        (PNG, "photo.txt"),  # binary data with a .txt name
        (make_zip(), "archive.zip"),  # a zip that is not a Word file
        (PNG, "photo.png"),
        (b"plain text", "notes.rtf"),  # text, but not .txt / .md
        (b"plain text", "notes"),  # no extension
    ],
    ids=["binary .txt", "zip", "png", "rtf", "no extension"],
)
def test_t5_unsupported_types(data, path):
    error = error_of(data, path)
    assert error.code is DocumentErrorCode.UNSUPPORTED_TYPE
    assert error.name == "Brief"


def test_t5_empty_text_file_has_no_text():
    assert error_of(b"  \n\n ", "empty.txt").code is DocumentErrorCode.NO_TEXT


def test_too_large_is_checked_first():
    error = error_of(b"x" * 11, "big.txt", max_bytes=10)
    assert error.code is DocumentErrorCode.TOO_LARGE


# --- T6: the content decides, not the name ---


def test_t6_pdf_without_extension_is_found():
    data = make_pdf([BRIEF_LINES])
    assert detect_kind(data, "") is DocumentKind.PDF
    assert BRIEF_LINES[0] in read(data, "renamed-file")


def test_t6_docx_with_a_wrong_extension_is_found():
    assert detect_kind(make_docx(), ".txt") is DocumentKind.DOCX


# --- helpers ---


@pytest.mark.parametrize(
    ("path", "extension"),
    [
        ("docs/Brief.TXT", ".txt"),
        ("a\\b\\c.md", ".md"),
        ("noext", ""),
        ("x.tar.gz", ".gz"),
    ],
)
def test_extension_of(path, extension):
    assert extension_of(path) == extension


def test_normalize_text():
    assert normalize_text("a\r\n\r\n\r\n\r\nb  \rc\x00") == "a\n\nb\nc"
